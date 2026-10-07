# -*- coding: utf-8 -*-
"""把项目 docs 里的 Markdown 转成可直接发送的 .docx。

为什么自己写：本机没有 pandoc；Git Bash 传中文路径给 python 容易乱码，
所以待转换清单写在脚本内部常量里，不在命令行传中文参数。

用法：
    venv\\Scripts\\python.exe tools\\md2docx.py
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(BASE, "docs")
OUT = os.path.join(DOCS, "待发送")

# (源 md 文件名, 输出 docx 文件名)
TASKS = [
    ("回复队长_第二轮反馈.md", "回复队长_第二轮反馈.docx"),
    ("给4号_院校表字段清单.md", "给4号_院校表字段清单.docx"),
]

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"
GREY = RGBColor(0x77, 0x77, 0x77)


def _set_font(run, name, size, bold=False, color=None, mono=False):
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), MONO_FONT if mono else name)
    rfonts.set(qn("w:hAnsi"), MONO_FONT if mono else name)
    rfonts.set(qn("w:eastAsia"), name)


def _add_text(par, text, size=10.5, bold=False, color=None, mono=False):
    """按 **粗体** 标记切分，粗体部分真的加粗。"""
    text = text.replace("`", "")
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            _set_font(run, BODY_FONT, size, bold=True, color=color, mono=mono)
        else:
            run = par.add_run(part)
            _set_font(run, BODY_FONT, size, bold=bold, color=color, mono=mono)


def _shade(par, hex_fill):
    pr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def _bottom_border(par):
    pr = par._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "BFBFBF")
    borders.append(bottom)
    pr.append(borders)


def _add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=len(rows), cols=len(header))
    table.style = "Table Grid"
    for c, cell_text in enumerate(header):
        cell = table.cell(0, c)
        cell.text = ""
        par = cell.paragraphs[0]
        _add_text(par, cell_text.replace("**", ""), size=10, bold=True)
        _shade(par, "EDEDED")
    for r, row in enumerate(body, start=1):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            par = cell.paragraphs[0]
            _add_text(par, cell_text, size=10)
    return table


def convert(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1

        # 代码块
        if line.strip().startswith("```"):
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # 吃掉收尾的 ```
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(12)
            par.paragraph_format.space_after = Pt(6)
            _shade(par, "F5F5F5")
            run = par.add_run("\n".join(block))
            _set_font(run, BODY_FONT, 9, mono=True)
            continue

        # 分隔线
        if line.strip() in ("---", "***", "___"):
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(2)
            _bottom_border(par)
            continue

        # 表格
        if line.strip().startswith("|"):
            rows, buf = [], []
            while i - 1 < len(lines) and lines[i - 1].strip().startswith("|"):
                raw = lines[i - 1].strip()
                cells = [c.strip() for c in raw.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    buf.append(cells)
                i += 1
            if buf:
                _add_table(doc, buf)
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2)
            sizes = {1: 16, 2: 14, 3: 12, 4: 11, 5: 10.5, 6: 10.5}
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
            par.paragraph_format.space_after = Pt(6)
            _add_text(par, text, size=sizes[level], bold=True)
            for run in par.runs:
                run.font.name = HEAD_FONT
                run._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
            continue

        # 引用
        if line.strip().startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(18)
            _add_text(par, line.strip().lstrip("> ").strip(), size=10, color=GREY)
            continue

        # 列表
        m = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m:
            par = doc.add_paragraph(style="List Bullet")
            _add_text(par, m.group(1), size=10.5)
            continue
        m = re.match(r"^\s*(\d+)[.)]\s+(.*)$", line)
        if m:
            par = doc.add_paragraph(style="List Number")
            _add_text(par, m.group(2), size=10.5)
            continue

        # 空行 / 普通段落
        if not line.strip():
            continue
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        _add_text(par, line.strip(), size=10.5)

    doc.save(docx_path)
    return docx_path


def main():
    os.makedirs(OUT, exist_ok=True)
    for src_name, dst_name in TASKS:
        src = os.path.join(DOCS, src_name)
        dst = os.path.join(OUT, dst_name)
        if not os.path.exists(src):
            print("找不到源文档：%s" % src)
            continue
        convert(src, dst)
        size_kb = os.path.getsize(dst) // 1024
        print("已生成：%s （%d KB）" % (dst, size_kb))


if __name__ == "__main__":
    sys.exit(main())
