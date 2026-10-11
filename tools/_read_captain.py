# -*- coding: utf-8 -*-
"""临时脚本：读队长回复 docx + 列出交付包 zip 内容（只列不解压）。"""
import os
import zipfile

from docx import Document

DOCX = r"D:\微信文档\xwechat_files\wxid_732wgfr20rjk22_c7f2\msg\file\2026-10\对二号第二轮反馈的组长回复_v2.docx"
ZIP = r"D:\微信文档\xwechat_files\wxid_732wgfr20rjk22_c7f2\msg\file\2026-10\组长任务交付包_v2_20261009.zip"

print("===== 队长回复 docx =====")
if os.path.exists(DOCX):
    doc = Document(DOCX)
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            print(t)
    for ti, table in enumerate(doc.tables):
        print("\n--- 表格 %d ---" % (ti + 1))
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            print(" | ".join(cells))
else:
    print("找不到 docx：%s" % DOCX)

print("\n===== 交付包 zip 内容 =====")
if os.path.exists(ZIP):
    z = zipfile.ZipFile(ZIP)
    for info in z.infolist():
        flag_utf8 = "UTF8" if info.flag_bits & 0x800 else "CP437?"
        print("%8d  %s  [%s]" % (info.file_size, info.filename, flag_utf8))
else:
    print("找不到 zip：%s" % ZIP)
